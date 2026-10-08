"""
unit_test/test_template_power_content.py
==========================================
BUG-17/18 (test_reports/2026-10-08_round4/BAO_CAO_TEST_VONG_4.md, còn mở
tới vòng 5 test_reports/2026-10-08_round5/BAO_CAO_TEST_VONG_5.md) — lỗi
NỘI DUNG trong chính file mẫu Word THẬT templates/TEMPLATE_POWER/*.docx
(không phải code Python):

  BUG-17: bienban.docx đề tiêu đề "BIÊN BẢN HIỆU CHUẨN" nhưng mục con lại
          ghi "KẾT QUẢ KIỂM ĐỊNH" — thuật ngữ "kiểm định"/"hiệu chuẩn" là 2
          khái niệm khác nhau trong đo lường, không được lẫn trong CÙNG 1
          văn bản.
  BUG-18: gcnkd.docx có dòng "TP Hồ Chí Minh, ngày   tháng   năm 20" để
          TRỐNG tay (không phải biến Jinja) — trong khi bienban.docx CÙNG
          mẫu đã dùng đúng {{ header.sign_date }} ở vị trí tương tự.

Test render qua core.report_templates.get_template() TRÊN FILE MẪU THẬT
(không phải fixture registry) để bắt được lỗi nằm NGAY TRONG file .docx.
"""

from datetime import date

import pytest

from core.session import CalibrationSession
from core.report_templates import get_template

pytest.importorskip("docx")
from docx import Document


def _render(tmp_path):
    session = CalibrationSession(template_id="TEMPLATE_POWER")
    session.meta.date = date(2026, 10, 8)
    tpl = get_template("TEMPLATE_POWER")
    bienban = tpl.generate_bienban(session, str(tmp_path / "bienban.docx"))
    gcnkd = tpl.generate_gcnkd(session, str(tmp_path / "gcnkd.docx"))
    return Document(bienban), Document(gcnkd)


def test_bienban_power_does_not_mix_kiem_dinh_with_hieu_chuan_title(tmp_path):
    bienban, _ = _render(tmp_path)
    full_text = "\n".join(p.text for p in bienban.paragraphs)
    assert "BIÊN BẢN HIỆU CHUẨN" in full_text
    assert "KIỂM ĐỊNH" not in full_text, (
        "bienban.docx (TEMPLATE_POWER) vẫn lẫn thuật ngữ 'kiểm định' trong "
        "văn bản 'hiệu chuẩn' (BUG-17)")
    assert "KẾT QUẢ HIỆU CHUẨN" in full_text


def test_gcnkd_power_sign_date_line_is_filled_not_blank(tmp_path):
    _, gcnkd = _render(tmp_path)
    full_text = "\n".join(p.text for p in gcnkd.paragraphs)
    assert "ngày      tháng      năm 20" not in full_text, (
        "Dòng ngày ký trên GCN (TEMPLATE_POWER) vẫn để trống tay (BUG-18)")
    assert "ngày 08 tháng 10 năm 2026" in full_text
