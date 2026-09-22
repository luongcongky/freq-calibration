"""
unit_test/test_xlsx_table_engine.py
======================================
Test core/xlsx_table_engine.py::render_xlsx_with_table_contexts() — song
song với unit_test/test_report_templates_generic.py nhưng cho mẫu Excel:
gõ text report_val('<id>')/result('<id>') vào ô, quét + ghi số liệu THÔ vào
đúng ô, giữ nguyên công thức khác trong sheet.
"""

from datetime import date

import openpyxl
import pytest

from core import table_engine
from core import xlsx_table_engine
from core.table_descriptor import TableDescriptor, RowDef
from core.scenario_runner import StepResult
from core.session import CalibrationSession, SessionMeta, DUTInfo, SessionTest


def _descriptor() -> TableDescriptor:
    return TableDescriptor(
        schema_version=1, table_id="T1", name="Bảng thử nghiệm", order=1,
        scenario_file="", layout="repeated_rows", value_unit="dBm",
        value_format="dbm_no_unit",
        rows=[
            RowDef(key="10 MHz", raw_count=3),
            RowDef(key="50 MHz", raw_count=1),
        ],
        columns=[], pass_rule={"type": "none"}, merge=[],
    )


def _session_with_result(descriptor: TableDescriptor) -> CalibrationSession:
    test = SessionTest(table_id="T1", name="Bảng thử nghiệm", enabled=True)
    test.step_results = [
        StepResult(action="report_val", ok=True, value=1.0),
        StepResult(action="report_val", ok=True, value=2.0),
        StepResult(action="report_val", ok=True, value=3.0),
        StepResult(action="report_val", ok=True, value=10.0),
    ]
    test.result_table = table_engine.map_table(descriptor, test.step_results)
    for r in test.result_table.rows:
        r.confirmed = True
    session = CalibrationSession(
        template_id="TEST", meta=SessionMeta(dut=DUTInfo(serial="SN1"), date=date(2026, 9, 17)),
        tests=[test],
    )
    return session


def _build_template_xlsx(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A2"] = "10 MHz"
    ws["B2"] = "report_val('T1')"
    ws["C2"] = "report_val('T1')"
    ws["D2"] = "report_val('T1')"
    ws["E2"] = "=AVERAGE(B2:D2)"
    ws["A3"] = "50 MHz"
    ws["B3"] = "report_val('T1')"
    ws["F2"] = "result('T1')"
    ws["G2"] = "Nhãn tĩnh không đụng tới"
    wb.save(str(path))


def test_render_xlsx_writes_raw_values_and_keeps_formulas(tmp_path):
    descriptor = _descriptor()
    session = _session_with_result(descriptor)

    tpl_path = tmp_path / "tpl.xlsx"
    _build_template_xlsx(tpl_path)
    out_path = tmp_path / "out.xlsx"

    xlsx_table_engine.render_xlsx_with_table_contexts(session, [descriptor], tpl_path, out_path)

    wb = openpyxl.load_workbook(str(out_path), data_only=False)
    ws = wb.active

    # report_val() -> SỐ THỰC (không phải chuỗi định dạng), đúng thứ tự
    assert ws["B2"].value == 1.0
    assert ws["C2"].value == 2.0
    assert ws["D2"].value == 3.0
    assert ws["B3"].value == 10.0

    # Công thức khác giữ NGUYÊN, không bị đụng tới
    assert ws["E2"].value == "=AVERAGE(B2:D2)"
    assert ws["G2"].value == "Nhãn tĩnh không đụng tới"

    # result('T1') — pass_rule "none" -> Đạt/Không đạt rỗng (openpyxl chuẩn
    # hoá chuỗi rỗng thành None khi lưu/đọc lại, xem test dưới)
    assert ws["F2"].value is None

    # Ép Excel tự tính lại công thức khi mở file kết quả
    assert wb.calculation.fullCalcOnLoad is True


def test_render_xlsx_missing_table_id_writes_empty_result(tmp_path):
    descriptor = _descriptor()
    session = _session_with_result(descriptor)

    tpl_path = tmp_path / "tpl.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "result('KHONG_TON_TAI')"
    wb.save(str(tpl_path))
    out_path = tmp_path / "out.xlsx"

    xlsx_table_engine.render_xlsx_with_table_contexts(session, [descriptor], tpl_path, out_path)

    wb2 = openpyxl.load_workbook(str(out_path))
    assert wb2.active["A1"].value is None


def test_build_raw_rows_by_table_returns_raw_readings(tmp_path):
    descriptor = _descriptor()
    session = _session_with_result(descriptor)

    raw_by_table = table_engine.build_raw_rows_by_table(session, [descriptor])
    assert list(raw_by_table.keys()) == ["T1"]
    rows = raw_by_table["T1"]
    assert [r.raw_readings for r in rows] == [[1.0, 2.0, 3.0], [10.0]]
