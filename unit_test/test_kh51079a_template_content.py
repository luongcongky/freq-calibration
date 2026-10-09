"""
unit_test/test_kh51079a_template_content.py
==============================================
N-01 (test_reports/2026-10-09_khong_thiet_bi_retest/BAO_CAO_TEST_LAI_K01_K12.md)
— lỗi NỘI DUNG trong chính file mẫu Excel THẬT
templates/KH_51079A_XLSX/bienban.xlsx (không phải code Python): ô B15
("Kết luận (chọn):") gõ CỨNG chữ "Đạt", không lấy từ ô Kết luận của app —
biên bản luôn ra "Kết luận (chọn): Đạt" dù mọi bài đo đều LỖI/chưa có số
liệu nào (sau khi K08 đã sửa app để ô Kết luận để TRỐNG trong trường hợp
đó). Sửa: B15 đổi sang marker header('conclusion') — đúng quy ước
core/xlsx_table_engine.py đã dùng cho B17/E17 trong CÙNG file.
"""

from datetime import date
from pathlib import Path
import zipfile

import pytest

from core.session import CalibrationSession
from core.report_templates import get_template

pytest.importorskip("openpyxl")
import openpyxl

_TEMPLATE_PATH = Path("templates/KH_51079A_XLSX/bienban.xlsx")

# R-03 (test_reports/2026-10-09_retest2/BAO_CAO_TEST_LAI_2.md): sửa N-01 lúc
# đầu dùng wb.load_workbook()+wb.save() của openpyxl để đổi B15 — openpyxl
# KHÔNG đọc/giữ được textbox "Dấu đơn vị" (vmlDrawing — đối tượng vẽ kiểu
# VML cũ) và ảnh logo trong HEADER TRANG (cũng qua vmlDrawing, khác ảnh
# trong THÂN sheet mà openpyxl giữ được) -> file 2,78 MB tụt xuống 32 KB,
# mất cả 2 đối tượng. Phải vá TRỰC TIẾP XML trong zip (core/xlsx_table_engine
# ._write_patched_xlsx — xem docstring đầu file đó), KHÔNG đi qua
# openpyxl.save() bao giờ cho file mẫu CÓ đối tượng vẽ/ảnh header.
_EXPECTED_ZIP_ENTRIES = {
    "xl/drawings/drawing1.xml", "xl/drawings/vmlDrawing1.vml",
    "xl/drawings/vmlDrawing2.vml", "xl/drawings/vmlDrawing3.vml",
    "xl/media/image1.png", "xl/media/image2.png",
    "xl/printerSettings/printerSettings1.bin",
    "xl/printerSettings/printerSettings2.bin",
    "xl/printerSettings/printerSettings3.bin",
    "xl/calcChain.xml",
}


def test_template_file_keeps_drawing_image_and_printer_objects():
    with zipfile.ZipFile(_TEMPLATE_PATH) as z:
        names = set(z.namelist())
    missing = _EXPECTED_ZIP_ENTRIES - names
    assert not missing, (
        f"Mất đối tượng (drawing/ảnh/printer) trong mẫu Excel — có thể do "
        f"sửa bằng openpyxl.save() (R-03): thiếu {missing}")


def test_template_file_size_matches_original_not_openpyxl_resave():
    size = _TEMPLATE_PATH.stat().st_size
    assert size > 1_000_000, (
        f"File mẫu chỉ còn {size} byte — nghi bị openpyxl.save() ghi lại mất "
        f"đối tượng (R-03, trước đây tụt từ 2,78 MB xuống 32 KB)")


def _render(tmp_path, conclusion: str):
    session = CalibrationSession(template_id="KH_51079A_XLSX")
    session.meta.date = date(2026, 10, 9)
    session.meta.conclusion = conclusion
    tpl = get_template("KH_51079A_XLSX")
    out = tpl.generate_bienban(session, str(tmp_path / "bienban.xlsx"))
    return openpyxl.load_workbook(out)


def test_template_source_no_longer_hardcodes_dat():
    """File mẫu THẬT trên đĩa không còn gõ cứng 'Đạt' ở B15."""
    wb = openpyxl.load_workbook("templates/KH_51079A_XLSX/bienban.xlsx")
    ws = wb["Thông tin chung"]
    assert ws["B15"].value != "Đạt", (
        "templates/KH_51079A_XLSX/bienban.xlsx vẫn gõ cứng 'Đạt' ở B15 (N-01)")
    assert "header(" in str(ws["B15"].value), (
        f"B15 không dùng marker header(...) như B17/E17 cùng file: {ws['B15'].value!r}")


def test_rendered_conclusion_blank_when_session_conclusion_empty(tmp_path):
    """Mọi bài LỖI, chưa có số liệu -> app để trống ô Kết luận (K08) -> biên
    bản xuất ra KHÔNG được tự ghi 'Đạt' (N-01)."""
    wb = _render(tmp_path, conclusion="")
    ws = wb["Thông tin chung"]
    assert ws["B15"].value in (None, ""), (
        f"Biên bản vẫn ghi kết luận dù app để trống (N-01): {ws['B15'].value!r}")


def test_rendered_conclusion_matches_session_when_set(tmp_path):
    """Không được vô tình đổi hành vi khi THỰC SỰ có kết luận."""
    wb = _render(tmp_path, conclusion="Đạt yêu cầu kỹ thuật đo lường")
    ws = wb["Thông tin chung"]
    assert ws["B15"].value == "Đạt yêu cầu kỹ thuật đo lường"
