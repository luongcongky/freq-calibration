"""
unit_test/test_xlsx_wizard_io.py
====================================
Test phần đọc Excel cụ thể của "Đọc bảng từ Excel"
(gui/template_manager_dialog.py::ImportTableFromExcelDialog gọi vào đây) —
tương đương unit_test/test_table_wizard_io.py nhưng cho .xlsx. Mọi hàm
format-agnostic (guess_column_role, build_rows_from_grid...) đã test sẵn ở
test_table_wizard_io.py, KHÔNG lặp lại ở đây.

QUYẾT ĐỊNH THIẾT KẾ (khách hàng chốt): khác bản Word, file .xlsx KHÔNG bao
giờ bị app ghi/sửa gì — mọi hàm ở đây CHỈ ĐỌC. Test nào cần 1 ô "đã có tag"
thì tự gán trực tiếp qua openpyxl (mô phỏng khách tự gõ tay trong Excel),
KHÔNG có hàm insert_report_val_tags() nào trong core/xlsx_wizard_io.py để gọi.
"""

import openpyxl

from core import xlsx_wizard_io as xwio


def _build_workbook(path, merge_summary_row: bool = True):
    """1 sheet: dòng tiêu đề + 1 dòng dữ liệu thường (3 cột đo riêng biệt) +
    1 dòng "Trung bình" GỘP 3 cột đo thành 1 ô rộng (mô phỏng dòng tổng hợp
    của khách — chỉ cần 1 report_val() dù có 3 cột "★ Giá trị đo" được chọn)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "Tần số"
    ws["B1"] = "V1"
    ws["C1"] = "V2"
    ws["D1"] = "V3"
    ws["E1"] = "TB"
    ws["A2"] = "10 MHz"
    ws["E2"] = "=AVERAGE(B2:D2)"
    ws["A3"] = "Trung bình"
    if merge_summary_row:
        ws.merge_cells("B3:D3")
    ws["E3"] = "=AVERAGE(B3:D3)"
    wb.save(str(path))
    return ws.title


def _tag_cells(path, sheet_name: str, coords: list, table_id: str) -> None:
    """Mô phỏng khách TỰ GÕ TAY report_val('<table_id>') vào các ô (coords:
    list[(row, col)] tuyệt đối, 1-based) NGAY TRONG EXCEL — dùng làm setup
    cho test, KHÔNG phải hàm sản phẩm (app không có hàm ghi tag nào)."""
    wb = openpyxl.load_workbook(str(path), data_only=False)
    ws = wb[sheet_name]
    for r, c in coords:
        ws.cell(row=r, column=c).value = f"report_val('{table_id}')"
    wb.save(str(path))


def test_list_sheets_reports_used_range_and_tag_state(tmp_path):
    path = tmp_path / "wb.xlsx"
    _build_workbook(path)

    sheets = xwio.list_sheets(path)
    assert len(sheets) == 1
    assert sheets[0].sheet_name == "Sheet1"
    assert sheets[0].already_tagged is False
    assert sheets[0].used_range.startswith("A1:")


def test_read_range_returns_text_grid(tmp_path):
    path = tmp_path / "wb.xlsx"
    _build_workbook(path)

    grid = xwio.read_range(path, "Sheet1", "A1:E2")
    assert grid[0] == ["Tần số", "V1", "V2", "V3", "TB"]
    assert grid[1][0] == "10 MHz"
    assert grid[1][4] == "=AVERAGE(B2:D2)"
    # Ô chưa có dữ liệu -> chuỗi rỗng, không phải None
    assert grid[1][1] == ""


# ---------------------------------------------------------------------------
# raw_counts_for_measured_cols — CHỈ ĐẾM ô đã có sẵn ĐÚNG text
# report_val('<table_id>') (khách tự gõ tay trong Excel), KHÔNG đếm ô rỗng
# hay ô có chữ/số khác — quyết định thiết kế đã chốt với khách hàng (tránh
# app tự ý ghi đè lên file bảng tính của khách).
# ---------------------------------------------------------------------------

def test_raw_counts_only_counts_cells_already_tagged(tmp_path):
    path = tmp_path / "wb.xlsx"
    _build_workbook(path)
    # Dòng 2: chỉ B2, C2 đã có tag — D2 vẫn rỗng (khách chưa gõ tới).
    _tag_cells(path, "Sheet1", [(2, 2), (2, 3)], "A1")

    counts = xwio.raw_counts_for_measured_cols(path, "Sheet1", [2], [2, 3, 4], "A1")
    assert counts == [2]


def test_raw_counts_ignores_empty_and_static_text_cells(tmp_path):
    path = tmp_path / "wb.xlsx"
    _build_workbook(path)
    ws_path = path
    wb = openpyxl.load_workbook(str(ws_path))
    ws = wb["Sheet1"]
    ws["C2"] = "0.5"          # chữ tĩnh khác, KHÔNG phải tag
    wb.save(str(ws_path))
    _tag_cells(path, "Sheet1", [(2, 2)], "A1")   # chỉ B2 có tag; C2="0.5", D2 rỗng

    counts = xwio.raw_counts_for_measured_cols(path, "Sheet1", [2], [2, 3, 4], "A1")
    assert counts == [1]


def test_raw_counts_requires_exact_table_id_match(tmp_path):
    path = tmp_path / "wb.xlsx"
    _build_workbook(path)
    _tag_cells(path, "Sheet1", [(2, 2), (2, 3)], "A9")   # tag của bảng KHÁC

    counts = xwio.raw_counts_for_measured_cols(path, "Sheet1", [2], [2, 3, 4], "A1")
    assert counts == [0]


def test_raw_counts_dedups_merged_measured_cells(tmp_path):
    path = tmp_path / "wb.xlsx"
    _build_workbook(path)
    # Dòng 2 (thường): B2, C2, D2 đã tag riêng biệt -> đếm 3.
    # Dòng 3 (gộp B3:D3): CHỈ ô góc B3 (anchor) có thể mang giá trị -> đếm 1
    # dù measured_cols_abs chọn cả B, C, D.
    _tag_cells(path, "Sheet1", [(2, 2), (2, 3), (2, 4), (3, 2)], "A1")

    counts = xwio.raw_counts_for_measured_cols(path, "Sheet1", [2, 3], [2, 3, 4], "A1")
    assert counts == [3, 1]


def test_raw_counts_all_untagged_returns_zero(tmp_path):
    path = tmp_path / "wb.xlsx"
    _build_workbook(path)

    counts = xwio.raw_counts_for_measured_cols(path, "Sheet1", [2, 3], [2, 3, 4], "A1")
    assert counts == [0, 0]


def test_find_missing_table_ids(tmp_path):
    path = tmp_path / "wb.xlsx"
    _build_workbook(path)
    _tag_cells(path, "Sheet1", [(2, 2), (2, 3), (2, 4), (3, 2)], "A1")

    assert xwio.find_missing_table_ids(path, ["A1"]) == []
    assert xwio.find_missing_table_ids(path, ["A1", "A9"]) == ["A9"]


def test_list_sheets_detects_existing_tags(tmp_path):
    path = tmp_path / "wb.xlsx"
    _build_workbook(path)
    _tag_cells(path, "Sheet1", [(2, 2), (2, 3), (2, 4), (3, 2)], "A1")

    sheets = xwio.list_sheets(path)
    assert sheets[0].already_tagged is True


# ---------------------------------------------------------------------------
# find_xlsx_table_grid — Bước 2/3 + màn "Sửa bảng" dùng để hiện lại cấu trúc
# thật, tương đương core/table_wizard_io.py::find_docx_table_grid.
# ---------------------------------------------------------------------------

def test_find_xlsx_table_grid_expands_left_and_up(tmp_path):
    path = tmp_path / "wb.xlsx"
    _build_workbook(path)
    _tag_cells(path, "Sheet1", [(2, 2), (2, 3), (2, 4), (3, 2)], "A1")

    grid = xwio.find_xlsx_table_grid(path, "A1")
    # min_col=B (2) -> mở rộng trái ra A (1); min_row=2 -> mở rộng lên dòng
    # tiêu đề (1), dừng vì dòng 0 không tồn tại.
    assert grid == [
        ["Tần số", "V1", "V2", "V3"],
        ["10 MHz", "report_val('A1')", "report_val('A1')", "report_val('A1')"],
        ["Trung bình", "report_val('A1')", "", ""],
    ]


def test_find_xlsx_table_grid_missing_table_id_returns_none(tmp_path):
    path = tmp_path / "wb.xlsx"
    _build_workbook(path)
    _tag_cells(path, "Sheet1", [(2, 2), (2, 3), (2, 4), (3, 2)], "A1")

    assert xwio.find_xlsx_table_grid(path, "A9") is None


def test_find_xlsx_table_grid_missing_file_returns_none(tmp_path):
    assert xwio.find_xlsx_table_grid(tmp_path / "khong_ton_tai.xlsx", "A1") is None


def test_find_xlsx_table_grid_stops_expanding_at_another_tables_tag(tmp_path):
    """2 bảng xếp chồng CÙNG 1 sheet — mở rộng lên của bảng dưới (A1) phải
    dừng lại NGAY trước dòng có tag của bảng trên (A0), không lấn sang."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A3"] = "10 MHz"
    ws["B3"] = "report_val('A0')"
    # dòng 4 để trống (khoảng cách giữa 2 bảng)
    ws["A5"] = "Tần số"
    ws["B5"] = "V1"
    ws["A6"] = "10 MHz"
    ws["B6"] = "report_val('A1')"
    wb.save(str(path := tmp_path / "wb.xlsx"))

    grid = xwio.find_xlsx_table_grid(path, "A1")
    assert grid == [
        ["", ""],            # dòng 4 (trống) — mở rộng lên tới đây rồi dừng
        ["Tần số", "V1"],
        ["10 MHz", "report_val('A1')"],
    ]


# ---------------------------------------------------------------------------
# first_data_row_index / value_columns_from_grid — dùng bởi gui/report_preview.py
# ::_xlsx_grid_for() để hiện đúng tiêu đề cột thật (như "lần 1"..."lần 5")
# ở bảng rà soát Bước 2/3, thay vì tên chung "Lần N" của _build_generic.
# ---------------------------------------------------------------------------

def test_first_data_row_index_finds_row_right_after_header(tmp_path):
    path = tmp_path / "wb.xlsx"
    _build_workbook(path)
    _tag_cells(path, "Sheet1", [(2, 2), (2, 3), (2, 4), (3, 2)], "A1")
    grid = xwio.find_xlsx_table_grid(path, "A1")
    # grid trả về từ find_xlsx_table_grid (test trên) = [header, "10 MHz" (dữ
    # liệu), "Trung bình"] -> dòng dữ liệu đầu tiên là index 1 (0-based).
    assert xwio.first_data_row_index(grid, "A1") == 1


def test_first_data_row_index_no_match_returns_zero():
    grid = [["Khoá", "V1"], ["10 MHz", "0.5"]]
    assert xwio.first_data_row_index(grid, "A1") == 0


def test_value_columns_from_grid_detects_report_val_with_argument():
    """Khác Word (report_val() rỗng), Excel dùng report_val('<id>') có tham
    số — value_columns_from_grid() phải nhận diện được cú pháp này."""
    grid = [
        ["Khoá", "V1", "V2", "TB"],
        ["10 MHz", "report_val('A1')", "report_val('A1')", "=AVERAGE(B2:C2)"],
    ]
    assert xwio.value_columns_from_grid(grid, header_row=0) == [1, 2]


def test_value_columns_from_grid_empty_grid_returns_empty_list():
    assert xwio.value_columns_from_grid([]) == []
