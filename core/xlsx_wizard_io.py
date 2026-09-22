"""
core/xlsx_wizard_io.py
=========================
Phần "đọc ô Excel cụ thể" của "Đọc bảng từ Excel" (gui/
template_manager_dialog.py::ImportTableFromExcelDialog gọi vào đây) — tương
đương core/table_wizard_io.py nhưng cho .xlsx thay vì .docx.

QUYẾT ĐỊNH THIẾT KẾ (khách hàng chốt, áp dụng CHUNG cho cả Word — xem
core/table_wizard_io.py::raw_counts_for_measured_cols): file này KHÔNG ghi
gì vào .xlsx của khách cả — CHỈ NHẬN DIỆN ô nào ĐÃ CÓ SẴN text
report_val('<table_id>') do khách TỰ GÕ TAY trong Excel từ trước. Cột
"★ Giá trị đo" khách chọn ở màn hình chỉ dùng để khoanh vùng TÌM, không có
tác dụng tự động gán report_val() vào ô rỗng/ô có chữ khác — tránh rủi ro
ghi đè ngoài ý muốn lên 1 file bảng tính khách có thể đã tự xây công thức
phức tạp xung quanh.

Mọi hàm/hằng số THAO TÁC TRÊN grid/WizardTableSpec/TableDescriptor thuần
(guess_column_role, build_rows_from_grid, WizardRowSpec, WizardTableSpec,
validate_table_id_available, validate_rows, build_descriptor,
write_descriptor_json, resolve_value_format, FORMAT_LABELS_ALL,
PASS_RULE_CHOICES, COLUMN_ROLE_CHOICES, guess_bare_number) đã KHÔNG phụ
thuộc docx — dùng lại thẳng từ core.table_wizard_io, không viết lại ở đây.

Không phụ thuộc Qt → test độc lập với phần GUI.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from openpyxl.utils.cell import range_boundaries

_TAG_RE = re.compile(r"^(report_val|gcn_avg|gcn_error|gcn_limit|result)\(")

# ---------------------------------------------------------------------------
# Đọc file .xlsx khách CHƯA gắn tag (hoặc đã gắn 1 phần) — CHỈ ĐỌC, không
# sửa/ghi gì cả (xem docstring đầu file).
# ---------------------------------------------------------------------------

@dataclass
class DetectedXlsxSheet:
    sheet_name: str
    used_range: str          # vd "A1:W27" — gợi ý vùng mặc định, khách tự sửa
    already_tagged: bool     # True nếu sheet đã có chữ "report_val(" ở đâu đó


def list_sheets(xlsx_path) -> list:
    """Liệt kê mọi sheet trong workbook kèm vùng dùng thật (gợi ý điền sẵn ô
    'Vùng dữ liệu') — CHỈ ĐỌC."""
    wb = openpyxl.load_workbook(str(xlsx_path), data_only=False)
    result = []
    for ws in wb.worksheets:
        already_tagged = False
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and "report_val(" in cell.value:
                    already_tagged = True
                    break
            if already_tagged:
                break
        result.append(DetectedXlsxSheet(
            sheet_name=ws.title, used_range=ws.dimensions, already_tagged=already_tagged,
        ))
    wb.close()
    return result


def parse_range(range_str: str) -> tuple:
    """(min_col, min_row, max_col, max_row) — 1-based, đúng quy ước openpyxl
    — bọc range_boundaries() để gui/template_manager_dialog.py không cần tự
    import openpyxl."""
    return range_boundaries(range_str)


def read_range(xlsx_path, sheet_name: str, range_str: str) -> list:
    """Đọc 1 vùng ô (vd 'A2:W26') thành grid (list[list[str]]) — nguyên văn
    text từng ô ('' nếu None, kể cả ô bị "nuốt" bởi merge — openpyxl tự trả
    None cho các ô đó, KHÔNG cần dedup như bảng Word có gridSpan)."""
    min_col, min_row, max_col, max_row = range_boundaries(range_str)
    wb = openpyxl.load_workbook(str(xlsx_path), data_only=False)
    ws = wb[sheet_name]
    grid = []
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        grid.append(["" if cell.value is None else str(cell.value) for cell in row])
    wb.close()
    return grid


def _merged_anchor(ws, row: int, col: int) -> tuple:
    """Toạ độ (row, col) Ô GÓC TRÊN-TRÁI của vùng merge chứa (row, col), hoặc
    chính (row, col) nếu không nằm trong merge nào — dùng để gộp nhiều cột
    "★ Giá trị đo" đã chọn nhưng thực chất cùng 1 ô vật lý (vd 1 dòng tổng
    hợp gộp ô) thành ĐÚNG 1 report_val(), giống nguyên tắc gridSpan bên Word."""
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return (rng.min_row, rng.min_col)
    return (row, col)


def raw_counts_for_measured_cols(xlsx_path, sheet_name: str, data_rows_abs: list,
                                  measured_cols_abs: list, table_id: str) -> list:
    """[raw_count_dòng_1, raw_count_dòng_2, ...] — đếm số ô "giá trị đo" mỗi
    dòng ĐÃ CÓ SẴN đúng text report_val('<table_id>') (sau khi gộp các cột đã
    chọn — measured_cols_abs, chỉ số cột TUYỆT ĐỐI trong sheet — cùng nằm
    trong 1 vùng merge lại thành 1 ô vật lý qua _merged_anchor, giống nguyên
    tắc gridSpan bên Word). CHỈ ĐỌC, không ghi gì — ô đang rỗng hoặc có chữ/
    số khác (chưa đúng tag) KHÔNG được tính, khách phải tự gõ tay
    report_val('<table_id>') vào Excel trước (xem docstring đầu file)."""
    tag = f"report_val('{table_id}')"
    wb = openpyxl.load_workbook(str(xlsx_path), data_only=False)
    ws = wb[sheet_name]
    counts = []
    for r in data_rows_abs:
        anchors = {_merged_anchor(ws, r, c) for c in measured_cols_abs}
        n = sum(1 for (ar, ac) in anchors
                if isinstance(ws.cell(row=ar, column=ac).value, str)
                and ws.cell(row=ar, column=ac).value.strip() == tag)
        counts.append(n)
    wb.close()
    return counts


def find_missing_table_ids(xlsx_path, table_ids: list) -> list:
    """Đọc toàn bộ ô chuỗi của mọi sheet trong 1 file .xlsx đã gắn tag tay,
    trả về danh sách table_id KHÔNG tìm thấy `'<ID>'` (đúng cú pháp
    report_val('ID')/gcn_avg('ID')/...) ở đâu trong file — cảnh báo sớm lỗi
    gõ nhầm mã bảng, KHÔNG sửa/chặn gì."""
    wb = openpyxl.load_workbook(str(xlsx_path), data_only=False)
    parts = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    parts.append(cell.value)
    wb.close()
    full_text = "\n".join(parts)
    return [tid for tid in table_ids if f"'{tid}'" not in full_text]


# ---------------------------------------------------------------------------
# "Bước 2/3 xem lại cấu trúc thật" — tương đương find_docx_table_grid() bên
# core/table_wizard_io.py (dùng bởi gui/report_preview.py VÀ màn "Sửa bảng")
# nhưng cho .xlsx. CHỈ ĐỌC.
# ---------------------------------------------------------------------------

def find_xlsx_table_grid(xlsx_path, table_id: str):
    """Tìm vùng ô trong 1 file .xlsx đã gắn tag report_val()/result()/gcn_*()
    của ĐÚNG table_id -> dựng lại grid (list[list[str]]) THAM KHẢO cấu trúc
    thật đó — tương đương find_docx_table_grid() bên bản Word, CHỈ ĐỌC,
    không sửa gì. None nếu file không tồn tại hoặc không ô nào gắn tag
    table_id này.

    Khác Word (1 bảng = 1 đối tượng doc.tables lấy TRỌN VẸN được ngay vì
    python-docx biết rõ ranh giới bảng): Excel không có khái niệm "bảng vật
    lý" — grid trả về là vùng CHỮ NHẬT NHỎ NHẤT bao hết mọi ô đã gắn tag của
    table_id này, mở rộng thêm ĐÚNG 1 cột bên trái (thường là cột nhãn/tần
    số, không có tag) và tối đa 6 dòng phía trên (thường là tiêu đề bảng) —
    dừng mở rộng lên ngay khi gặp 1 dòng có ô gắn tag của MỘT table_id KHÁC,
    để không lấn sang bảng kế bên trong cùng sheet."""
    xlsx_path = Path(xlsx_path)
    if not xlsx_path.exists():
        return None
    needle = f"'{table_id}'"
    wb = openpyxl.load_workbook(str(xlsx_path), data_only=False)
    try:
        for ws in wb.worksheets:
            matches = [(cell.row, cell.column) for row in ws.iter_rows() for cell in row
                       if isinstance(cell.value, str) and needle in cell.value]
            if not matches:
                continue

            min_row = min(r for r, _ in matches)
            max_row = max(r for r, _ in matches)
            min_col = min(c for _, c in matches)
            max_col = max(c for _, c in matches)
            if min_col > 1:
                min_col -= 1

            def _row_has_other_tag(r: int) -> bool:
                for c in range(1, ws.max_column + 1):
                    v = ws.cell(row=r, column=c).value
                    if isinstance(v, str) and needle not in v and _TAG_RE.match(v.strip()):
                        return True
                return False

            top = min_row
            for _ in range(6):
                candidate = top - 1
                if candidate < 1 or _row_has_other_tag(candidate):
                    break
                top = candidate
            min_row = top

            return [["" if (v := ws.cell(row=r, column=c).value) is None else str(v)
                     for c in range(min_col, max_col + 1)]
                    for r in range(min_row, max_row + 1)]
        return None
    finally:
        wb.close()


def first_data_row_index(grid: list, table_id: str) -> int:
    """Chỉ số dòng (0-based, TƯƠNG ĐỐI trong `grid`) ĐẦU TIÊN có ít nhất 1 ô
    report_val('<table_id>') — dùng để suy ra dòng NGAY TRÊN nó là dòng tiêu
    đề thật (vd "lần 1"..."lần 5") khi dựng bảng rà soát Bước 2/3 (gui/
    report_preview.py::_xlsx_grid_for). Khác bảng Word (find_docx_table_grid
    đọc TRỌN 1 bảng vật lý, dòng 0 LUÔN là tiêu đề) — grid từ
    find_xlsx_table_grid() có thể có vài dòng trống/tiêu đề phụ phía TRÊN
    trước khi tới dòng dữ liệu đầu tiên, không đơn giản là dòng 0. Trả 0 nếu
    không tìm thấy (không nên xảy ra với grid lấy từ find_xlsx_table_grid)."""
    needle = f"'{table_id}'"
    for i, row in enumerate(grid):
        if any(isinstance(c, str) and needle in c for c in row):
            return i
    return 0


def value_columns_from_grid(grid: list, header_row: int = 0) -> list:
    """Tương đương core/table_wizard_io.py::value_columns_from_grid nhưng
    nhận diện cú pháp report_val('<id>') (CÓ tham số) thay vì report_val()
    rỗng của Word — trả danh sách chỉ số cột (0-based, trái→phải) có ít nhất
    1 ô report_val() ở 1 dòng khác header_row."""
    if not grid:
        return []
    n_cols = max(len(r) for r in grid)
    cols = []
    for c in range(n_cols):
        for r_i, row in enumerate(grid):
            if r_i == header_row:
                continue
            if c < len(row) and isinstance(row[c], str) and "report_val(" in row[c]:
                cols.append(c)
                break
    return cols
