"""
scripts/format_report_templates.py
====================================
Định dạng lại giao diện (font/cỡ chữ/đậm/nghiêng/căn lề/viền bảng/tô nền)
cho các file mẫu Biên Bản/GCN (.docx) của TEMPLATE_FREQ/TEMPLATE_POWER —
KHÔNG đổi 1 ký tự nội dung nào, kể cả tag Jinja {{ }}/{% %}/{%tr %} (đã
verify bằng cách so sánh text trích xuất trước/sau phải giống hệt, và
render thử bằng docxtpl phải thành công, không còn tag thô sót lại).

Phong cách: đơn sắc cổ điển (đen/xám, không màu) — Times New Roman, phân
cấp bằng đậm/nghiêng/cỡ chữ, bảng có viền mảnh + tiêu đề bảng tô nền xám
nhạt, theo đúng lựa chọn của khách hàng khi được hỏi (không dùng xanh navy
hay bất kỳ màu sắc nào khác, giữ tinh thần văn bản hành chính nhà nước).

Cách phân loại (áp dụng theo thứ tự, dừng ở điều kiện khớp đầu tiên):
  - Đoạn văn: khớp đúng 1 trong các câu tiêu đề đã biết ("BIÊN BẢN KIỂM
    ĐỊNH"...) -> tiêu đề; chỉ có tag điều khiển {% if %}/{% endif %} ->
    tô xám nhỏ (không hiện trong file xuất thật, chỉ để admin dễ nhận khi
    sửa mẫu); bắt đầu bằng tên cơ quan (letterhead gộp vào 1 đoạn, không
    phải bảng) -> letterhead; khớp mẫu "<địa danh>, ngày ..." -> dòng
    ngày ký, căn phải nghiêng; "Bảng Ax - ..." -> chú thích bảng; số thứ
    tự "1 ...", "3.1 ...", "A.2 ..." -> tiêu đề phụ; ngoặc đơn ngắn -> phụ
    đề; có "{{" -> dòng nhãn (in đậm phần nhãn trước mỗi tag, giữ
    thường phần tag); còn lại -> văn bản thường.
  - Bảng: có "CỘNG HÒA" ở ô bên phải, 1 dòng 2 cột -> letterhead dạng
    bảng; có "Trang"/"No of paper" -> bảng chân trang; có
    "report_val()"/".result" -> bảng dữ liệu (dòng đầu tô nền xám nhạt +
    đậm + viền, dòng {%tr if%}/{%tr endif%} không viền vì không hiện
    trong file xuất thật); còn lại -> bảng chữ ký (đậm chức danh, giữ
    thường tên/tag, căn giữa).

Dùng:
    python scripts/format_report_templates.py

Sau khi chạy, PHẢI kiểm tra lại bằng:
    pytest unit_test/ -q
và xem trước file .docx thật (bằng Word/LibreOffice) trước khi coi là
hoàn tất — kiểm thị giác không thể tự động hoá hoàn toàn.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Times New Roman"
SZ_BODY = Pt(13)
SZ_TABLE = Pt(12)
SZ_TITLE = Pt(16)
SZ_SUBTITLE = Pt(12)
SZ_SECTION = Pt(13)
SZ_SMALL = Pt(11)
SZ_LETTERHEAD = Pt(9.5)
GRAY = RGBColor(0x59, 0x59, 0x59)
LIGHT_GRAY_FILL = "F2F2F2"

TAG_RE = re.compile(r"(\{\{[^}]*\}\}|\{%\s*tr\s+if[^%]*%\}|\{%\s*tr\s+endif\s*%\}"
                     r"|\{%\s*if[^%]*%\}|\{%\s*endif\s*%\})")

TITLE_TEXTS = {
    "BIÊN BẢN KIỂM ĐỊNH", "BIÊN BẢN HIỆU CHUẨN",
    "GIẤY CHỨNG NHẬN KIỂM ĐỊNH", "GIẤY CHỨNG NHẬN HIỆU CHUẨN",
}
SECTION_TEXTS = {"KẾT QUẢ KIỂM ĐỊNH", "KẾT QUẢ HIỆU CHUẨN"}
SUBHEADER_RE = re.compile(r"^(\d+(\.\d+)*|[A-Z]\.\d+(\.\d+)*)[\s.]")  # "1 Kiểm tra...", "3.1 Xác định...", "A.3.1 ..."
CAPTION_RE = re.compile(r"^Bảng [A-Z0-9]+ *-")
# Chỉ khớp {% if %}/{% endif %}/{%tr ...%} — KHÔNG khớp {{ value }} (tag giá
# trị VẪN hiện nội dung thật sau khi render, không được tô xám/thu nhỏ như
# tag điều khiển {% %} — lỗi đã gặp với dòng {{ header.today }} đứng riêng).
CONTROL_ONLY_RE = re.compile(r"^\{%.*%\}$")
LETTERHEAD_PREFIXES = ("CỤC TIÊU CHUẨN ĐO LƯỜNG", "CỤC TL-ĐL-CL")
SIGN_DATE_LINE_RE = re.compile(r"^[^,]+,\s*ngày\s")
# Dòng chỉ có đúng {{ header.sign_date }} (đồng nghiệp đã đổi từ chữ ghi
# cứng "TP Hồ Chí Minh, ngày..." sang tag động) — vẫn là dòng "địa danh,
# ngày ký", chỉ khác là giờ nội dung đến từ tag thay vì chữ tĩnh.
BARE_SIGN_DATE_TAG_RE = re.compile(r"^\{\{\s*header\.sign_date\s*\}\}$")


def _set_font(run, size=None, bold=None, italic=None, color=None, name=FONT):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), name)
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def _style_all_runs(paragraph, **kw):
    for r in paragraph.runs:
        _set_font(r, **kw)


def _split_bold_labels(paragraph, value_size=SZ_BODY):
    """Tách PHẦN NHÃN (chữ tĩnh trước mỗi tag {{ }}) ra đậm, phần tag/giá
    trị giữ thường — áp dụng cho các dòng "Nhãn: {{ tag }}". Không đổi 1 ký
    tự nào của text, chỉ ghi lại thành nhiều run với định dạng khác nhau.
    Chỉ xử lý khi đoạn có ĐÚNG 1 run gốc (mọi mẫu hiện có đều vậy — được
    viết tay trực tiếp, chưa qua docxtpl render nên chưa bị tách run)."""
    if len(paragraph.runs) != 1:
        return False
    original = paragraph.runs[0]
    text = original.text
    if "{{" not in text and "{%" not in text:
        return False
    pieces = TAG_RE.split(text)
    if len(pieces) <= 1:
        return False
    p_elem = paragraph._p
    for r in list(paragraph.runs):
        p_elem.remove(r._element)
    for piece in pieces:
        if not piece:
            continue
        is_tag = bool(TAG_RE.fullmatch(piece))
        run = paragraph.add_run(piece)
        _set_font(run, size=value_size, bold=False if is_tag else True)
    return True


def style_title(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(4)
    _style_all_runs(paragraph, size=SZ_TITLE, bold=True)


def style_subtitle(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(8)
    _style_all_runs(paragraph, size=SZ_SUBTITLE, italic=True)


def style_section(paragraph):
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(4)
    _style_all_runs(paragraph, size=SZ_SECTION, bold=True)


def style_subheader(paragraph):
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(2)
    _style_all_runs(paragraph, size=SZ_BODY, bold=True)


def style_caption(paragraph):
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(2)
    _style_all_runs(paragraph, size=SZ_BODY, italic=True, bold=True)


def style_jinja_control(paragraph):
    """{% if %}/{% endif %} đứng riêng 1 đoạn — KHÔNG hiện trong file xuất
    thật (docxtpl tự xoá khi render), chỉ tô xám nhạt cho dễ phân biệt khi
    admin mở FILE MẪU ra sửa."""
    _style_all_runs(paragraph, size=Pt(9), color=GRAY)


def style_label_line(paragraph):
    if not _split_bold_labels(paragraph):
        _style_all_runs(paragraph, size=SZ_BODY)


def style_plain(paragraph, size=SZ_BODY):
    _style_all_runs(paragraph, size=size)


def style_letterhead_paragraph(paragraph):
    """Vài mẫu (GCN) gộp cả khối letterhead vào 1 đoạn văn thay vì bảng 2
    cột như Biên Bản — giữ NGUYÊN cấu trúc đó (không chuyển thành bảng,
    tránh đổi cấu trúc nội dung), chỉ thu nhỏ cỡ chữ + căn giữa cho đúng
    tinh thần 1 khối letterhead."""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(6)
    _style_all_runs(paragraph, size=SZ_SMALL)


def style_sign_date_line(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _style_all_runs(paragraph, size=SZ_BODY, italic=True)


def classify_and_style_paragraph(paragraph):
    text = paragraph.text.strip()
    if not text:
        return
    if text in TITLE_TEXTS:
        style_title(paragraph)
    elif text in SECTION_TEXTS:
        style_section(paragraph)
    elif CONTROL_ONLY_RE.match(text):
        style_jinja_control(paragraph)
    elif text.startswith(LETTERHEAD_PREFIXES):
        style_letterhead_paragraph(paragraph)
    elif (SIGN_DATE_LINE_RE.match(text) and "{{" not in text) or BARE_SIGN_DATE_TAG_RE.match(text):
        style_sign_date_line(paragraph)
    elif CAPTION_RE.match(text):
        style_caption(paragraph)
    elif SUBHEADER_RE.match(text) and "{{" not in text:
        style_subheader(paragraph)
    elif text.startswith("(") and text.endswith(")") and "{{" not in text and len(text) < 40:
        style_subtitle(paragraph)
    elif "{{" in text:
        style_label_line(paragraph)
    else:
        style_plain(paragraph)


# ---------------------------------------------------------------------------
# Bảng
# ---------------------------------------------------------------------------

def _set_cell_border(cell, sz=4, color="000000"):
    tcpr = cell._tc.get_or_add_tcPr()
    old = tcpr.find(qn('w:tcBorders'))
    if old is not None:
        tcpr.remove(old)
    borders = OxmlElement('w:tcBorders')
    for edge in ("top", "start", "bottom", "end"):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), str(sz))
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)
        borders.append(el)
    tcpr.append(borders)


def _shade_cell(cell, hex_color):
    tcpr = cell._tc.get_or_add_tcPr()
    old = tcpr.find(qn('w:shd'))
    if old is not None:
        tcpr.remove(old)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcpr.append(shd)


def _cell_center_v(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    old = tcpr.find(qn('w:vAlign'))
    if old is not None:
        tcpr.remove(old)
    va = OxmlElement('w:vAlign')
    va.set(qn('w:val'), 'center')
    tcpr.append(va)


def _style_cell_text(cell, size=SZ_TABLE, bold=None, center=True):
    for p in cell.paragraphs:
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            _set_font(r, size=size, bold=bold)


def is_letterhead_table(table):
    return len(table.rows) == 1 and len(table.columns) == 2 and "CỘNG HÒA" in table.cell(0, 1).text


def is_footer_table(table):
    t = table.rows[0].cells[0].text
    return len(table.rows) == 1 and ("Trang" in t or "No of paper" in t)


def is_data_table(table):
    full = " ".join(c.text for row in table.rows for c in row.cells)
    return "report_val()" in full or ".result" in full


def style_letterhead_table(table):
    left, right = table.rows[0].cells
    for p in left.paragraphs:
        for r in p.runs:
            _set_font(r, size=SZ_SMALL)
    for p in right.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            _set_font(r, size=SZ_LETTERHEAD, bold=True)


def style_footer_table(table):
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    _set_font(r, size=Pt(10), italic=True)


def style_signature_table(table):
    for row in table.rows:
        for cell in row.cells:
            _cell_center_v(cell)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(2)
                for r in p.runs:
                    is_tag = bool(TAG_RE.fullmatch(r.text.strip())) if r.text.strip() else False
                    _set_font(r, size=SZ_BODY, bold=not is_tag)


def style_data_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_row = table.rows[0]
    for cell in header_row.cells:
        _shade_cell(cell, LIGHT_GRAY_FILL)
        _cell_center_v(cell)
        _style_cell_text(cell, size=SZ_TABLE, bold=True, center=True)
        _set_cell_border(cell)
    for row in table.rows[1:]:
        full_row_text = " ".join(c.text for c in row.cells)
        is_control_row = bool(CONTROL_ONLY_RE.match(full_row_text.strip()))
        for cell in row.cells:
            _cell_center_v(cell)
            # Cột đầu (nhãn dòng, vd "10 MHz") thường là text -> giữ center
            # cho đồng nhất với các cột số liệu, KHÔNG đổi text.
            _style_cell_text(cell, size=SZ_TABLE, bold=False, center=True)
            if not is_control_row:
                _set_cell_border(cell)


def style_table(table):
    if is_letterhead_table(table):
        style_letterhead_table(table)
    elif is_footer_table(table):
        style_footer_table(table)
    elif is_data_table(table):
        style_data_table(table)
    else:
        style_signature_table(table)


def format_document(in_path, out_path=None) -> None:
    out_path = out_path or in_path
    doc = Document(in_path)
    # Font mặc định toàn tài liệu (lưới an toàn — mọi run đã được set tường
    # minh ở trên rồi, đây chỉ là fallback cho phần không có run nào, vd
    # paragraph mark trống).
    normal = doc.styles['Normal']
    normal.font.name = FONT
    normal.font.size = SZ_BODY
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), FONT)

    for p in doc.paragraphs:
        classify_and_style_paragraph(p)
    for t in doc.tables:
        style_table(t)
    # Bảng lồng trong ô (hiện tại các mẫu không có, nhưng duyệt cho chắc).
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for nested in cell.tables:
                    style_table(nested)

    doc.save(out_path)


TEMPLATES_DIR = Path(__file__).resolve().parent.parent / "templates"
DEFAULT_TARGETS = [
    TEMPLATES_DIR / "TEMPLATE_FREQ" / "bienban.docx",
    TEMPLATES_DIR / "TEMPLATE_FREQ" / "gcnkd.docx",
    TEMPLATES_DIR / "TEMPLATE_POWER" / "bienban.docx",
    TEMPLATES_DIR / "TEMPLATE_POWER" / "gcnkd.docx",
]


def main():
    for path in DEFAULT_TARGETS:
        if not path.exists():
            print(f"BỎ QUA (không tồn tại): {path}")
            continue
        format_document(path)
        print(f"Đã định dạng lại: {path}")


if __name__ == "__main__":
    main()
